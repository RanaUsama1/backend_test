"""Excel view of an already-captured archive. Never performs network/database IO."""
import json
from io import BytesIO
from threading import BoundedSemaphore
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill, Alignment

MAX_BYTES = 4 * 1024 * 1024
MAX_ROWS = 30000
_gate = BoundedSemaphore(1)


def flatten(value, path='', depth=0):
    if depth > 48: raise ValueError('nesting_limit')
    if isinstance(value, dict):
        yield path, 'object', ''
        for k,v in value.items():
            yield from flatten(v,path+'/'+k.replace('~','~0').replace('/','~1'),depth+1)
    elif isinstance(value,list):
        yield path,'array',''
        for i,v in enumerate(value):yield from flatten(v,path+'/'+str(i),depth+1)
    else:
        kind='null' if value is None else 'boolean' if isinstance(value,bool) else 'number' if isinstance(value,(int,float)) else 'string'
        yield path,kind,json.dumps(value,ensure_ascii=True,allow_nan=False)


def build_archive_workbook(normalized, manifest, entries, database="assembly"):
    if not _gate.acquire(False):raise ValueError('excel_generation_busy')
    try:return _build(normalized,manifest,entries,database)
    finally:_gate.release()


def _build(normalized,manifest,entries,database):
    payload=json.dumps(normalized,ensure_ascii=True,allow_nan=False)
    if len(payload.encode()) + sum(len(b) for _,b in entries if b is not None)>MAX_BYTES:
        raise ValueError('excel_input_limit_4_MiB')
    plans=[];row_count=0
    def add(name, headers, rows):
        nonlocal row_count
        clean=[]
        for row in rows:
            if any(isinstance(v,str) and len(v)>30000 for v in row):raise ValueError('excel_cell_limit')
            clean.append(row);row_count+=1
            if row_count>MAX_ROWS:raise ValueError('excel_row_limit_30000')
        plans.append((name,headers,clean))
    def fields(data):
        for path,kind,value in flatten(data):
            chunks=[value[i:i+30000] for i in range(0,len(value),30000)] or ['']
            for i,c in enumerate(chunks,1):yield json.dumps(path,ensure_ascii=True),kind,i,len(chunks),c
    headings=['Path JSON','Type','Part','Parts','Value JSON']
    d=normalized if isinstance(normalized,dict) else {}
    summary=[('Accession',manifest.get('accession', d.get('tax_id', 'diagnostic'))),('Organism',d.get('organism_name')),('Assembly',d.get('assembly_name')),
             ('Genome size (bp)',d.get('genome_size_bp')),('Chromosomes',d.get('number_of_chromosomes')),
             ('Total genes',d.get('total_genes')),('Protein-coding genes',d.get('protein_coding_genes')),
             ('Retrieval status',manifest.get('retrieval_status', manifest.get('status'))),('Snapshot UTC',manifest.get('created_at', d.get('fetched_at'))),
             ('Scope','Captured metadata responses only; not every linked NCBI record.'),
             ('Reading detail sheets','Values are JSON text; null, false, zero and empty strings remain distinct. Join numbered parts before JSON decoding.'),
             ('Original Text','Each source body is UTF-8 decoded and JSON escaped, then split into ordered parts. Original ZIP files remain authoritative.'),
             ('Raw XML and FTP','Complete source text is in Original Text; no XML field mapping is needed to retain its contents.')]
    if database == 'taxonomy':
        summary = [(k, d.get(k)) for k in ('tax_id','requested_tax_id','query','scientific_name','common_name','rank','division','parent_tax_id','fetched_at','status')]
        summary += [('Scope', manifest.get('scope')), ('Details', 'All normalized fields and original source text appear on the remaining sheets.')]
        add('Taxonomy Summary',['Field','Value'],summary)
        add('Lineage',['Tax ID','Scientific name','Rank'],[(r.get('tax_id'),r.get('name'),r.get('rank')) for r in d.get('lineage',[])])
        add('Genetic Codes',['Code type','ID','Name'],[(k,(d.get(k) or {}).get('id'),(d.get(k) or {}).get('name')) for k in ('genetic_code','mitochondrial_genetic_code')])
    elif database != 'assembly':
        summary = [('Database',database), ('Record',manifest.get('requested_accession') or manifest.get('gene_id') or d.get('accession')), ('Export mode',manifest.get('excel_export',{}).get('mode')), ('Scope','Only captured metadata; see Manifest for Excel exclusions and retrieval limits.')]
        summary += [(k,v) for k,v in d.items() if not isinstance(v,(dict,list)) and len(str(v)) < 2000][:25]
        add('Record Summary',['Field','Value'],summary)
    else:
        add('Assembly Summary',['Field','Value'],summary)
    add('Normalized Metadata',headings,fields(normalized))
    add('Manifest',headings,fields(manifest))
    originals=[]
    for index,((_,body),source) in enumerate(zip(entries,manifest['sources']),1):
        if body is None:continue
        # Strict decoding: fail visibly instead of replacing unsupported bytes.
        text=body.decode('utf-8')
        encoded=json.dumps(text,ensure_ascii=True)
        chunks=[encoded[i:i+30000] for i in range(0,len(encoded),30000)] or ['']
        for i,c in enumerate(chunks,1):originals.append((source['file'],i,len(chunks),c))
        if 'json' in source.get('content_type','').lower():
            try:parsed=json.loads(text)
            except ValueError:continue # Unparseable JSON still retained as original text.
            add('Source '+str(index).zfill(3),headings,fields(parsed))
    add('Original Text',['Source file','Part','Parts','Text JSON'],originals)
    wb=Workbook(write_only=True)
    for name,headers,rows in plans:
        ws=wb.create_sheet(name);ws.freeze_panes='A2'
        for i in range(len(headers)):ws.column_dimensions[chr(65+i)].width=65 if i==0 else 95 if i==len(headers)-1 else 14
        for row_index,row in enumerate([headers]+rows):
            cells=[]
            for value in row:
                cell=WriteOnlyCell(ws,value=value)
                if isinstance(value,str):cell.data_type='s'
                if row_index==0:
                    cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='17365D')
                if name in ('Assembly Summary','Taxonomy Summary'):cell.alignment=Alignment(wrap_text=True,vertical='top')
                cells.append(cell)
            ws.append(cells)
        ws.auto_filter.ref='A1:'+chr(64+len(headers))+str(len(rows)+1)
    stream=BytesIO();wb.save(stream);return stream.getvalue()
