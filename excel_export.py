"""Bounded export of a client-supplied snapshot. No database or network calls."""
import json
import re
from io import BytesIO
from threading import BoundedSemaphore
from datetime import datetime, timezone
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill

MAX_BYTES = 2 * 1024 * 1024
MAX_ROWS = 20000
MAX_DEPTH = 48
_gate = BoundedSemaphore(1)  # Per worker, not global across Render workers.


def field_rows(value, path='', depth=0):
    if depth > MAX_DEPTH:
        raise ValueError('Metadata nesting exceeds the Excel export limit.')
    if isinstance(value, dict):
        yield path, 'object', ''
        for key, child in value.items():
            escaped = key.replace('~', '~0').replace('/', '~1')
            yield from field_rows(child, path + '/' + escaped, depth + 1)
    elif isinstance(value, list):
        yield path, 'array', ''
        for i, child in enumerate(value):
            yield from field_rows(child, path + '/' + str(i), depth + 1)
    else:
        kind = ('null' if value is None else 'boolean' if isinstance(value, bool)
                else 'number' if isinstance(value, (int, float)) else 'string')
        yield path, kind, json.dumps(value, ensure_ascii=True, allow_nan=False)


def build_workbook(database, record_id, metadata):
    # Validate before starting the writer; never silently truncate data.
    rows = []
    for path, kind, value in field_rows(metadata):
        path = json.dumps(path, ensure_ascii=True)  # reversible control-character escaping
        if len(path) > 30000:
            raise ValueError('A field path exceeds the Excel export limit.')
        parts = [value[i:i+30000] for i in range(0, len(value), 30000)] or ['']
        for i, part in enumerate(parts, 1):
            rows.append((path, kind, i, len(parts), part))
            if len(rows) > MAX_ROWS:
                raise ValueError('Metadata exceeds 20,000 Excel rows. Use the original-source ZIP.')
    wb = Workbook(write_only=True)
    def sheet(name, headings, data, widths):
        ws = wb.create_sheet(name)
        ws.freeze_panes = 'A2'
        for col, width in widths.items(): ws.column_dimensions[col].width = width
        def cells(values, header=False):
            result=[]
            for value in values:
                c=WriteOnlyCell(ws, value=value)
                if isinstance(value, str):c.data_type='s'  # Never execute supplied text as a formula.
                if header:
                    c.font=Font(bold=True,color='FFFFFF')
                    c.fill=PatternFill('solid',fgColor='17365D')
                result.append(c)
            return result
        ws.append(cells(headings, True))
        count=1
        for row in data:ws.append(cells(row));count+=1
        ws.auto_filter.ref='A1:'+chr(64+len(headings))+str(count)
    sheet('Overview', ['Field','Value'], [
        ('Database',database),('Record label',record_id),
        ('Exported UTC',datetime.now(timezone.utc).isoformat()),
        ('Scope','Client-supplied search snapshot; not independently verified or freshly fetched.'),
        ('Coverage','Every submitted JSON field is included. Fields absent from this snapshot cannot be recovered here.'),
        ('Reading values','Metadata values use JSON text to preserve null, false, zero, empty strings and numeric precision.'),
        ('Long values','Join Value JSON parts in Part order before decoding JSON.'),
        ('Paths','Decode Path JSON as JSON, then interpret as an RFC 6901 JSON Pointer. Root is an empty pointer.'),
        ('Original sources','Keep the existing source ZIP for original response bytes and omitted records.')
    ], {'A':25,'B':110})
    sheet('Metadata', ['Path JSON','Type','Part','Parts','Value JSON'], rows,
          {'A':65,'B':14,'C':9,'D':9,'E':100})
    out=BytesIO();wb.save(out);out.seek(0)
    return out


def register_excel_export(app):
    from flask import request, jsonify, send_file
    @app.post('/api/export/excel')
    def export_excel():
        if not _gate.acquire(blocking=False):
            return jsonify(message='An Excel export is in progress. Please retry shortly.'),429
        try:
            if request.content_length is not None and request.content_length > MAX_BYTES:
                return jsonify(message='Excel export is limited to 2 MiB. Use the source ZIP for larger records.'),413
            if not request.is_json:
                return jsonify(message='Send application/json.'),415
            # Bounded read also covers requests without a Content-Length header.
            body=request.stream.read(MAX_BYTES+1)
            if len(body)>MAX_BYTES:
                return jsonify(message='Excel export is limited to 2 MiB.'),413
            def reject_constant(value):raise ValueError('Non-finite numbers are not supported.')
            payload=json.loads(body, parse_constant=reject_constant)
            if not isinstance(payload,dict):raise ValueError('Expected a JSON object.')
            database=payload.get('database')
            if database not in ('assembly','gene','protein','nucleotide','taxonomy'):
                raise ValueError('Unsupported database.')
            label=payload.get('record_id')
            if not isinstance(label,str) or not re.fullmatch(r'[A-Za-z0-9_.-]{1,100}',label):
                raise ValueError('Provide a valid record identifier.')
            metadata=payload.get('metadata')
            if not isinstance(metadata,dict) or not metadata:
                raise ValueError('A nonempty metadata object is required.')
            stream=build_workbook(database,label,metadata)
            response=send_file(stream, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                               as_attachment=True, download_name=database+'_'+label+'.xlsx',max_age=0)
            response.headers['Cache-Control']='no-store'
            response.call_on_close(stream.close)
            return response
        except (ValueError, RecursionError, UnicodeError):
            return jsonify(message='Invalid metadata or export limit exceeded. Use the original-source ZIP for larger records.'),400
        finally:
            _gate.release()
    return export_excel
